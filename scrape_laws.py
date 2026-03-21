#!/usr/bin/env python3
"""
سكريبت لجلب الأحكام القضائية المنشورة من البوابة القانونية لوزارة العدل السعودية
https://laws.moj.gov.sa/

يدعم الجلب التدريجي (incremental) مع حفظ التقدم واستئناف العمل.
يصدّر النتائج إلى ملف Word (.docx).
"""

import json
import os
import re
import sys
import time
import logging
from pathlib import Path
from datetime import datetime

import requests
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ─── الإعدادات ───────────────────────────────────────────────────────────────

BASE_URL = "https://laws.moj.gov.sa"
LIST_URL = f"{BASE_URL}/ar/JudicialDecisionsList"
PROGRESS_FILE = "scrape_progress.json"
OUTPUT_DOCX = "احكام_وزارة_العدل.docx"
REQUEST_DELAY = 2  # ثوانٍ بين كل طلب (لتجنب الحظر)
MAX_RETRIES = 3
RETRY_DELAY = 5  # ثوانٍ قبل إعادة المحاولة

# إعداد التسجيل
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler("scrape_log.txt", encoding="utf-8"),
    ],
)
logger = logging.getLogger(__name__)

# ─── Headers لمحاكاة المتصفح ─────────────────────────────────────────────────

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "ar,en;q=0.9",
    "Accept-Encoding": "gzip, deflate, br",
    "Connection": "keep-alive",
    "Referer": f"{BASE_URL}/ar",
}


# ─── إدارة التقدم ────────────────────────────────────────────────────────────


def load_progress():
    """تحميل ملف التقدم المحفوظ"""
    if os.path.exists(PROGRESS_FILE):
        with open(PROGRESS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {
        "last_page": 0,
        "total_pages": None,
        "scraped_urls": [],
        "decisions_count": 0,
        "last_updated": None,
    }


def save_progress(progress):
    """حفظ التقدم"""
    progress["last_updated"] = datetime.now().isoformat()
    with open(PROGRESS_FILE, "w", encoding="utf-8") as f:
        json.dump(progress, f, ensure_ascii=False, indent=2)


# ─── جلب البيانات ─────────────────────────────────────────────────────────────


def create_session():
    """إنشاء جلسة requests مع إعدادات مناسبة"""
    session = requests.Session()
    session.headers.update(HEADERS)
    # زيارة الصفحة الرئيسية أولاً للحصول على الكوكيز
    try:
        session.get(f"{BASE_URL}/ar", timeout=30)
        time.sleep(1)
    except Exception as e:
        logger.warning(f"لم يتم الوصول للصفحة الرئيسية: {e}")
    return session


def fetch_page(session, url, retries=MAX_RETRIES):
    """جلب صفحة مع إعادة المحاولة"""
    for attempt in range(retries):
        try:
            response = session.get(url, timeout=30)
            response.raise_for_status()
            return response.text
        except requests.RequestException as e:
            logger.warning(
                f"محاولة {attempt + 1}/{retries} فشلت لـ {url}: {e}"
            )
            if attempt < retries - 1:
                wait = RETRY_DELAY * (attempt + 1)
                logger.info(f"انتظار {wait} ثوانٍ قبل إعادة المحاولة...")
                time.sleep(wait)
    return None


def get_total_pages(session):
    """معرفة العدد الإجمالي للصفحات"""
    html = fetch_page(session, f"{LIST_URL}/1")
    if not html:
        return None

    soup = BeautifulSoup(html, "lxml")

    # البحث عن عناصر الترقيم (pagination)
    pagination = soup.select(".pagination a, .pager a, [class*='page'] a")
    if pagination:
        page_numbers = []
        for link in pagination:
            text = link.get_text(strip=True)
            if text.isdigit():
                page_numbers.append(int(text))
            # التحقق من href أيضاً
            href = link.get("href", "")
            match = re.search(r"/(\d+)(?:\?|$)", href)
            if match:
                page_numbers.append(int(match.group(1)))
        if page_numbers:
            return max(page_numbers)

    # البحث عن نص يحتوي على العدد الكلي
    text = soup.get_text()
    total_match = re.search(r"(\d+)\s*(?:صفحة|من|إجمالي|total|pages)", text)
    if total_match:
        return int(total_match.group(1))

    # محاولة البحث عن الصفحة الأخيرة في الـ pagination
    last_page_link = soup.select_one(
        'a[aria-label="Last"], a[title="الأخيرة"], .pagination li:last-child a'
    )
    if last_page_link:
        href = last_page_link.get("href", "")
        match = re.search(r"/(\d+)", href)
        if match:
            return int(match.group(1))

    logger.warning("لم يتم تحديد عدد الصفحات - سيتم الجلب حتى عدم وجود نتائج")
    return None


def extract_decision_links(html):
    """استخراج روابط الأحكام من صفحة القائمة"""
    soup = BeautifulSoup(html, "lxml")
    links = []

    # البحث عن روابط الأحكام
    for a in soup.find_all("a", href=True):
        href = a["href"]
        if "JudicialDecisionsList/0/" in href or "JudicialDecision/" in href:
            full_url = href if href.startswith("http") else f"{BASE_URL}{href}"
            if full_url not in links:
                links.append(full_url)

    return links


def extract_decision_content(html, url):
    """استخراج محتوى حكم قضائي من صفحته"""
    soup = BeautifulSoup(html, "lxml")
    decision = {
        "url": url,
        "title": "",
        "court": "",
        "date": "",
        "number": "",
        "category": "",
        "content": "",
        "details": {},
    }

    # استخراج العنوان
    title_el = soup.find("h1") or soup.find("h2") or soup.find(class_=re.compile(r"title"))
    if title_el:
        decision["title"] = title_el.get_text(strip=True)

    # استخراج التفاصيل من الجداول أو divs
    # البحث عن أزواج العنوان/القيمة
    detail_patterns = [
        (r"المحكمة|court", "court"),
        (r"التاريخ|date|تاريخ", "date"),
        (r"الرقم|number|رقم", "number"),
        (r"التصنيف|category|نوع", "category"),
    ]

    # محاولة استخراج من جدول
    for table in soup.find_all("table"):
        for row in table.find_all("tr"):
            cells = row.find_all(["td", "th"])
            if len(cells) >= 2:
                label = cells[0].get_text(strip=True)
                value = cells[1].get_text(strip=True)
                decision["details"][label] = value
                for pattern, key in detail_patterns:
                    if re.search(pattern, label, re.IGNORECASE):
                        decision[key] = value

    # محاولة استخراج من dl/dt/dd
    for dl in soup.find_all("dl"):
        dts = dl.find_all("dt")
        dds = dl.find_all("dd")
        for dt, dd in zip(dts, dds):
            label = dt.get_text(strip=True)
            value = dd.get_text(strip=True)
            decision["details"][label] = value
            for pattern, key in detail_patterns:
                if re.search(pattern, label, re.IGNORECASE):
                    decision[key] = value

    # محاولة استخراج من divs مع labels
    for div in soup.find_all("div", class_=re.compile(r"detail|info|field|row")):
        label_el = div.find(class_=re.compile(r"label|key|title"))
        value_el = div.find(class_=re.compile(r"value|data|content"))
        if label_el and value_el:
            label = label_el.get_text(strip=True)
            value = value_el.get_text(strip=True)
            decision["details"][label] = value
            for pattern, key in detail_patterns:
                if re.search(pattern, label, re.IGNORECASE):
                    decision[key] = value

    # استخراج المحتوى الرئيسي
    content_el = soup.find(
        class_=re.compile(r"content|body|text|article|decision")
    ) or soup.find("article")
    if content_el:
        # إزالة العناصر غير المرغوب فيها
        for unwanted in content_el.find_all(
            ["nav", "header", "footer", "script", "style"]
        ):
            unwanted.decompose()
        decision["content"] = content_el.get_text("\n", strip=True)
    else:
        # محاولة استخراج من main أو div رئيسي
        main = soup.find("main") or soup.find(id=re.compile(r"content|main"))
        if main:
            for unwanted in main.find_all(
                ["nav", "header", "footer", "script", "style"]
            ):
                unwanted.decompose()
            decision["content"] = main.get_text("\n", strip=True)

    return decision


# ─── تصدير Word ──────────────────────────────────────────────────────────────


def create_or_load_document():
    """إنشاء أو تحميل ملف Word"""
    if os.path.exists(OUTPUT_DOCX):
        logger.info(f"تحميل الملف الموجود: {OUTPUT_DOCX}")
        return Document(OUTPUT_DOCX)

    doc = Document()

    # إعداد النمط الأساسي
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Arial"
    font.size = Pt(12)
    font.rtl = True

    # العنوان الرئيسي
    title = doc.add_heading("الأحكام القضائية المنشورة", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph(
        f"البوابة القانونية - وزارة العدل السعودية\n"
        f"تاريخ الجلب: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("─" * 50)

    return doc


def add_decision_to_doc(doc, decision, index):
    """إضافة حكم قضائي إلى ملف Word"""

    # عنوان الحكم
    title_text = decision.get("title", f"حكم رقم {index}")
    heading = doc.add_heading(f"{index}. {title_text}", level=1)

    # التفاصيل
    details_items = []
    if decision.get("court"):
        details_items.append(f"المحكمة: {decision['court']}")
    if decision.get("number"):
        details_items.append(f"رقم الحكم: {decision['number']}")
    if decision.get("date"):
        details_items.append(f"التاريخ: {decision['date']}")
    if decision.get("category"):
        details_items.append(f"التصنيف: {decision['category']}")

    # إضافة أي تفاصيل إضافية
    for label, value in decision.get("details", {}).items():
        if value and label not in [
            decision.get("court"),
            decision.get("number"),
            decision.get("date"),
            decision.get("category"),
        ]:
            details_items.append(f"{label}: {value}")

    if details_items:
        details_para = doc.add_paragraph()
        for item in details_items:
            run = details_para.add_run(f"{item}\n")
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(100, 100, 100)

    # المحتوى
    content = decision.get("content", "")
    if content:
        # تقسيم المحتوى إلى فقرات
        paragraphs = content.split("\n")
        for para_text in paragraphs:
            para_text = para_text.strip()
            if para_text:
                p = doc.add_paragraph(para_text)
                p.paragraph_format.space_after = Pt(6)

    # خط فاصل
    doc.add_paragraph("─" * 50)

    return doc


def save_document(doc):
    """حفظ ملف Word"""
    doc.save(OUTPUT_DOCX)
    logger.info(f"تم حفظ الملف: {OUTPUT_DOCX}")


# ─── الجلب الرئيسي ────────────────────────────────────────────────────────────


def scrape_decisions():
    """الدالة الرئيسية للجلب التدريجي"""
    progress = load_progress()
    logger.info("=" * 60)
    logger.info("بدء جلب الأحكام القضائية من البوابة القانونية")
    logger.info(f"التقدم المحفوظ: صفحة {progress['last_page']}, "
                f"{progress['decisions_count']} حكم")
    logger.info("=" * 60)

    session = create_session()

    # معرفة عدد الصفحات
    if progress["total_pages"] is None:
        logger.info("جاري تحديد عدد الصفحات...")
        total = get_total_pages(session)
        if total:
            progress["total_pages"] = total
            logger.info(f"عدد الصفحات: {total}")
        else:
            logger.info("لم يتم تحديد العدد - سيتم الجلب حتى نفاد النتائج")
            progress["total_pages"] = 999  # حد أقصى افتراضي

    # تحميل أو إنشاء ملف Word
    doc = create_or_load_document()

    # بدء الجلب من حيث توقفنا
    start_page = progress["last_page"] + 1
    decision_index = progress["decisions_count"]

    page = start_page
    consecutive_empty = 0
    max_empty = 3  # عدد الصفحات الفارغة قبل التوقف

    try:
        while page <= progress["total_pages"] and consecutive_empty < max_empty:
            logger.info(f"\n📄 جلب الصفحة {page}...")
            html = fetch_page(session, f"{LIST_URL}/{page}")

            if not html:
                logger.error(f"فشل جلب الصفحة {page}")
                consecutive_empty += 1
                page += 1
                continue

            # استخراج روابط الأحكام
            links = extract_decision_links(html)
            logger.info(f"   وُجد {len(links)} حكم في الصفحة {page}")

            if not links:
                consecutive_empty += 1
                page += 1
                time.sleep(REQUEST_DELAY)
                continue

            consecutive_empty = 0

            # جلب كل حكم
            for link in links:
                if link in progress["scraped_urls"]:
                    logger.info(f"   ⏭ تم جلبه مسبقاً: {link[:80]}...")
                    continue

                time.sleep(REQUEST_DELAY)
                logger.info(f"   📥 جلب حكم: {link[:80]}...")

                decision_html = fetch_page(session, link)
                if not decision_html:
                    logger.warning(f"   ⚠ فشل جلب الحكم: {link[:80]}")
                    continue

                # استخراج المحتوى
                decision = extract_decision_content(decision_html, link)
                decision_index += 1

                # إضافة إلى Word
                add_decision_to_doc(doc, decision, decision_index)

                # تحديث التقدم
                progress["scraped_urls"].append(link)
                progress["decisions_count"] = decision_index

                logger.info(
                    f"   ✅ حكم #{decision_index}: "
                    f"{decision.get('title', 'بدون عنوان')[:50]}"
                )

                # حفظ تدريجي كل 5 أحكام
                if decision_index % 5 == 0:
                    save_document(doc)
                    save_progress(progress)
                    logger.info(f"   💾 حفظ تدريجي ({decision_index} حكم)")

            # تحديث رقم الصفحة
            progress["last_page"] = page
            save_progress(progress)
            page += 1

            time.sleep(REQUEST_DELAY)

    except KeyboardInterrupt:
        logger.info("\n⚠ تم الإيقاف بواسطة المستخدم")
    except Exception as e:
        logger.error(f"خطأ غير متوقع: {e}", exc_info=True)
    finally:
        # حفظ نهائي
        save_document(doc)
        save_progress(progress)
        logger.info(f"\n{'=' * 60}")
        logger.info(f"تم حفظ {decision_index} حكم في {OUTPUT_DOCX}")
        logger.info(f"آخر صفحة: {progress['last_page']}")
        logger.info(f"يمكنك استئناف الجلب لاحقاً بتشغيل السكريبت مجدداً")
        logger.info(f"{'=' * 60}")


# ─── نقطة البداية ─────────────────────────────────────────────────────────────

if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(
        description="جلب الأحكام القضائية من البوابة القانونية لوزارة العدل"
    )
    parser.add_argument(
        "--reset", action="store_true", help="إعادة البدء من الصفر"
    )
    parser.add_argument(
        "--delay",
        type=float,
        default=REQUEST_DELAY,
        help=f"التأخير بين الطلبات بالثواني (افتراضي: {REQUEST_DELAY})",
    )
    parser.add_argument(
        "--output",
        type=str,
        default=OUTPUT_DOCX,
        help=f"اسم ملف الإخراج (افتراضي: {OUTPUT_DOCX})",
    )
    parser.add_argument(
        "--max-pages",
        type=int,
        default=None,
        help="الحد الأقصى لعدد الصفحات المراد جلبها",
    )

    args = parser.parse_args()

    if args.reset and os.path.exists(PROGRESS_FILE):
        os.remove(PROGRESS_FILE)
        logger.info("تم إعادة تعيين التقدم")

    REQUEST_DELAY = args.delay
    OUTPUT_DOCX = args.output

    if args.max_pages:
        progress = load_progress()
        progress["total_pages"] = args.max_pages
        save_progress(progress)

    scrape_decisions()
